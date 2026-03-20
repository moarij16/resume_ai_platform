"""
Resume document renderer.

Converts a GeneratedDocument (resume or refactor type) into a downloadable
PDF or DOCX file using a modern, ATS-friendly template.

The AI response JSON is expected to have:
  summary, experience, education, skills, improvements_made
"""

import io
import json

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    HRFlowable,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
)

from docx import Document as DocxDocument  # type: ignore[import-untyped]
from docx.enum.text import WD_ALIGN_PARAGRAPH  # type: ignore[import-untyped]
from docx.oxml.ns import qn  # type: ignore[import-untyped]
from docx.shared import Pt, RGBColor, Inches  # type: ignore[import-untyped]
from docx.oxml import OxmlElement  # type: ignore[import-untyped]

# ── Colour palette ────────────────────────────────────────────────────────────
ACCENT = colors.HexColor("#1A1A2E")  # deep navy — section headers / name
RULE = colors.HexColor("#3D5A80")  # steel blue — horizontal rules
MUTED = colors.HexColor("#555555")  # dark grey — meta text


# ── Helpers ───────────────────────────────────────────────────────────────────


def _parse_content(document_content: str) -> dict:
    """Return the parsed JSON dict from document_content."""
    try:
        return json.loads(document_content)
    except (json.JSONDecodeError, TypeError):
        return {}


def _full_name(user) -> str:
    parts = [user.first_name, user.last_name]
    name = " ".join(p for p in parts if p).strip()
    return name or user.email


# ── PDF renderer ──────────────────────────────────────────────────────────────


def render_pdf(document) -> bytes:
    """Render a GeneratedDocument as a PDF and return the raw bytes."""
    data = _parse_content(document.document_content)
    user = document.user
    buffer = io.BytesIO()

    doc = SimpleDocTemplate(
        buffer,
        pagesize=LETTER,
        leftMargin=0.75 * inch,
        rightMargin=0.75 * inch,
        topMargin=0.75 * inch,
        bottomMargin=0.75 * inch,
    )

    base = getSampleStyleSheet()

    # ── Custom styles ──────────────────────────────────────────────────────
    style_name = ParagraphStyle(
        "CandidateName",
        parent=base["Normal"],
        fontSize=22,
        fontName="Helvetica-Bold",
        textColor=ACCENT,
        alignment=TA_CENTER,
        spaceAfter=2,
    )
    style_contact = ParagraphStyle(
        "Contact",
        parent=base["Normal"],
        fontSize=9,
        fontName="Helvetica",
        textColor=MUTED,
        alignment=TA_CENTER,
        spaceAfter=4,
    )
    style_section = ParagraphStyle(
        "Section",
        parent=base["Normal"],
        fontSize=11,
        fontName="Helvetica-Bold",
        textColor=ACCENT,
        spaceBefore=10,
        spaceAfter=2,
    )
    style_body = ParagraphStyle(
        "Body",
        parent=base["Normal"],
        fontSize=9.5,
        fontName="Helvetica",
        leading=14,
        spaceAfter=3,
    )
    style_meta = ParagraphStyle(
        "Meta",
        parent=base["Normal"],
        fontSize=9,
        fontName="Helvetica-Oblique",
        textColor=MUTED,
        spaceAfter=2,
    )
    style_bullet = ParagraphStyle(
        "Bullet",
        parent=base["Normal"],
        fontSize=9.5,
        fontName="Helvetica",
        leading=14,
        leftIndent=14,
        bulletIndent=4,
        spaceAfter=1,
    )

    story = []

    # ── Header ─────────────────────────────────────────────────────────────
    story.append(Paragraph(_full_name(user), style_name))
    story.append(Paragraph(user.email, style_contact))
    story.append(HRFlowable(width="100%", thickness=1.5, color=RULE, spaceAfter=6))

    # ── Summary ────────────────────────────────────────────────────────────
    summary = data.get("summary", "")
    if summary:
        story.append(Paragraph("PROFESSIONAL SUMMARY", style_section))
        story.append(HRFlowable(width="100%", thickness=0.5, color=RULE, spaceAfter=4))
        story.append(Paragraph(summary, style_body))

    # ── Experience ─────────────────────────────────────────────────────────
    experience = data.get("experience", [])
    if experience:
        story.append(Paragraph("EXPERIENCE", style_section))
        story.append(HRFlowable(width="100%", thickness=0.5, color=RULE, spaceAfter=4))
        for job in experience:
            title = job.get("title", "")
            company = job.get("company", "")
            duration = job.get("duration", "")
            story.append(Paragraph(f"<b>{title}</b>  —  {company}", style_body))
            if duration:
                story.append(Paragraph(duration, style_meta))
            for bullet in job.get("bullets", []):
                story.append(Paragraph(f"• {bullet}", style_bullet))
            story.append(Spacer(1, 4))

    # ── Education ──────────────────────────────────────────────────────────
    education = data.get("education", [])
    if education:
        story.append(Paragraph("EDUCATION", style_section))
        story.append(HRFlowable(width="100%", thickness=0.5, color=RULE, spaceAfter=4))
        for edu in education:
            degree = edu.get("degree", "")
            institution = edu.get("institution", "")
            year = edu.get("year", "")
            line = f"<b>{degree}</b>  —  {institution}"
            if year:
                line += f"  ({year})"
            story.append(Paragraph(line, style_body))

    # ── Skills ─────────────────────────────────────────────────────────────
    skills = data.get("skills", [])
    if skills:
        story.append(Paragraph("SKILLS", style_section))
        story.append(HRFlowable(width="100%", thickness=0.5, color=RULE, spaceAfter=4))
        story.append(Paragraph("  •  ".join(skills), style_body))

    doc.build(story)
    return buffer.getvalue()


# ── DOCX renderer ─────────────────────────────────────────────────────────────


def _add_rule(doc: DocxDocument) -> None:
    """Add a thin horizontal rule paragraph."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "3D5A80")
    pBdr.append(bottom)
    pPr.append(pBdr)


def _section_heading(doc: DocxDocument, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)
    _add_rule(doc)


def render_docx(document) -> bytes:
    """Render a GeneratedDocument as a DOCX and return the raw bytes."""
    data = _parse_content(document.document_content)
    user = document.user
    buffer = io.BytesIO()

    doc = DocxDocument()

    # ── Page margins ───────────────────────────────────────────────────────
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # ── Header ─────────────────────────────────────────────────────────────
    name_p = doc.add_paragraph()
    name_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name_p.paragraph_format.space_after = Pt(2)
    name_run = name_p.add_run(_full_name(user))
    name_run.bold = True
    name_run.font.size = Pt(22)
    name_run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)

    contact_p = doc.add_paragraph()
    contact_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact_p.paragraph_format.space_after = Pt(4)
    contact_run = contact_p.add_run(user.email)
    contact_run.font.size = Pt(9)
    contact_run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    _add_rule(doc)

    # ── Summary ────────────────────────────────────────────────────────────
    summary = data.get("summary", "")
    if summary:
        _section_heading(doc, "PROFESSIONAL SUMMARY")
        p = doc.add_paragraph(summary)
        p.paragraph_format.space_after = Pt(4)
        for run in p.runs:
            run.font.size = Pt(9.5)

    # ── Experience ─────────────────────────────────────────────────────────
    experience = data.get("experience", [])
    if experience:
        _section_heading(doc, "EXPERIENCE")
        for job in experience:
            title = job.get("title", "")
            company = job.get("company", "")
            duration = job.get("duration", "")

            title_p = doc.add_paragraph()
            title_p.paragraph_format.space_after = Pt(1)
            t_run = title_p.add_run(title)
            t_run.bold = True
            t_run.font.size = Pt(9.5)
            if company:
                title_p.add_run(f"  —  {company}").font.size = Pt(9.5)

            if duration:
                dur_p = doc.add_paragraph(duration)
                dur_p.paragraph_format.space_after = Pt(2)
                for r in dur_p.runs:
                    r.italic = True
                    r.font.size = Pt(9)
                    r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

            for bullet in job.get("bullets", []):
                bp = doc.add_paragraph(style="List Bullet")
                bp.paragraph_format.left_indent = Inches(0.25)
                bp.paragraph_format.space_after = Pt(1)
                bp_run = bp.add_run(bullet)
                bp_run.font.size = Pt(9.5)

            doc.add_paragraph().paragraph_format.space_after = Pt(2)

    # ── Education ──────────────────────────────────────────────────────────
    education = data.get("education", [])
    if education:
        _section_heading(doc, "EDUCATION")
        for edu in education:
            degree = edu.get("degree", "")
            institution = edu.get("institution", "")
            year = edu.get("year", "")
            ep = doc.add_paragraph()
            ep.paragraph_format.space_after = Pt(3)
            dr = ep.add_run(degree)
            dr.bold = True
            dr.font.size = Pt(9.5)
            suffix = f"  —  {institution}"
            if year:
                suffix += f"  ({year})"
            ep.add_run(suffix).font.size = Pt(9.5)

    # ── Skills ─────────────────────────────────────────────────────────────
    skills = data.get("skills", [])
    if skills:
        _section_heading(doc, "SKILLS")
        sp = doc.add_paragraph("  •  ".join(skills))
        sp.paragraph_format.space_after = Pt(4)
        for r in sp.runs:
            r.font.size = Pt(9.5)

    doc.save(buffer)
    return buffer.getvalue()
