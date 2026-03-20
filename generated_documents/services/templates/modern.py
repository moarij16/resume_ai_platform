"""
Modern resume template — clean, ATS-friendly, navy accent.

Name: "modern"
"""

import io

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import HRFlowable, Paragraph, SimpleDocTemplate, Spacer

from docx import Document as DocxDocument  # type: ignore[import-untyped]
from docx.enum.text import WD_ALIGN_PARAGRAPH  # type: ignore[import-untyped]
from docx.oxml import OxmlElement  # type: ignore[import-untyped]
from docx.oxml.ns import qn  # type: ignore[import-untyped]
from docx.shared import Inches, Pt, RGBColor  # type: ignore[import-untyped]

from .base import ResumeTemplate

# ── Colour palette ────────────────────────────────────────────────────────────
_ACCENT = colors.HexColor("#1A1A2E")  # deep navy
_RULE   = colors.HexColor("#3D5A80")  # steel blue
_MUTED  = colors.HexColor("#555555")  # dark grey


class ModernTemplate(ResumeTemplate):
    """Modern ATS-friendly template with navy accents and rule separators."""

    # ── PDF ───────────────────────────────────────────────────────────────────

    def render_pdf(self, document) -> bytes:
        data   = self.parse_content(document.document_content)
        user   = document.user
        buffer = io.BytesIO()

        doc = SimpleDocTemplate(
            buffer,
            pagesize=LETTER,
            leftMargin=0.75 * inch,
            rightMargin=0.75 * inch,
            topMargin=0.75 * inch,
            bottomMargin=0.75 * inch,
        )

        base = getSampleStyleSheet()

        s_name = ParagraphStyle(
            "CandidateName", parent=base["Normal"],
            fontSize=22, fontName="Helvetica-Bold",
            textColor=_ACCENT, alignment=TA_CENTER, spaceAfter=2,
        )
        s_contact = ParagraphStyle(
            "Contact", parent=base["Normal"],
            fontSize=9, fontName="Helvetica",
            textColor=_MUTED, alignment=TA_CENTER, spaceAfter=4,
        )
        s_section = ParagraphStyle(
            "Section", parent=base["Normal"],
            fontSize=11, fontName="Helvetica-Bold",
            textColor=_ACCENT, spaceBefore=10, spaceAfter=2,
        )
        s_body = ParagraphStyle(
            "Body", parent=base["Normal"],
            fontSize=9.5, fontName="Helvetica", leading=14, spaceAfter=3,
        )
        s_meta = ParagraphStyle(
            "Meta", parent=base["Normal"],
            fontSize=9, fontName="Helvetica-Oblique",
            textColor=_MUTED, spaceAfter=2,
        )
        s_bullet = ParagraphStyle(
            "Bullet", parent=base["Normal"],
            fontSize=9.5, fontName="Helvetica", leading=14,
            leftIndent=14, bulletIndent=4, spaceAfter=1,
        )

        def rule(thickness=0.5):
            return HRFlowable(width="100%", thickness=thickness, color=_RULE, spaceAfter=4)

        story = []

        # Header
        story.append(Paragraph(self.full_name(user), s_name))
        story.append(Paragraph(user.email, s_contact))
        story.append(HRFlowable(width="100%", thickness=1.5, color=_RULE, spaceAfter=6))

        # Summary
        summary = data.get("summary", "")
        if summary:
            story += [Paragraph("PROFESSIONAL SUMMARY", s_section), rule(), Paragraph(summary, s_body)]

        # Experience
        for job in data.get("experience", []):
            if not story or story[-1] != Paragraph("EXPERIENCE", s_section):
                story += [Paragraph("EXPERIENCE", s_section), rule()]
            story.append(Paragraph(f"<b>{job.get('title','')}</b>  —  {job.get('company','')}", s_body))
            if job.get("duration"):
                story.append(Paragraph(job["duration"], s_meta))
            for bullet in job.get("bullets", []):
                story.append(Paragraph(f"• {bullet}", s_bullet))
            story.append(Spacer(1, 4))

        # Education
        education = data.get("education", [])
        if education:
            story += [Paragraph("EDUCATION", s_section), rule()]
            for edu in education:
                line = f"<b>{edu.get('degree','')}</b>  —  {edu.get('institution','')}"
                if edu.get("year"):
                    line += f"  ({edu['year']})"
                story.append(Paragraph(line, s_body))

        # Skills
        skills = data.get("skills", [])
        if skills:
            story += [Paragraph("SKILLS", s_section), rule(),
                      Paragraph("  •  ".join(skills), s_body)]

        doc.build(story)
        return buffer.getvalue()

    # ── DOCX ──────────────────────────────────────────────────────────────────

    def render_docx(self, document) -> bytes:
        data   = self.parse_content(document.document_content)
        user   = document.user
        buffer = io.BytesIO()
        doc    = DocxDocument()

        for section in doc.sections:
            section.top_margin    = Inches(0.75)
            section.bottom_margin = Inches(0.75)
            section.left_margin   = Inches(0.75)
            section.right_margin  = Inches(0.75)

        # Header
        name_p = doc.add_paragraph()
        name_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        name_p.paragraph_format.space_after = Pt(2)
        nr = name_p.add_run(self.full_name(user))
        nr.bold = True
        nr.font.size = Pt(22)
        nr.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)

        contact_p = doc.add_paragraph()
        contact_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        contact_p.paragraph_format.space_after = Pt(4)
        cr = contact_p.add_run(user.email)
        cr.font.size = Pt(9)
        cr.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

        self._docx_rule(doc)

        # Summary
        summary = data.get("summary", "")
        if summary:
            self._docx_heading(doc, "PROFESSIONAL SUMMARY")
            p = doc.add_paragraph(summary)
            p.paragraph_format.space_after = Pt(4)
            for r in p.runs:
                r.font.size = Pt(9.5)

        # Experience
        experience = data.get("experience", [])
        if experience:
            self._docx_heading(doc, "EXPERIENCE")
            for job in experience:
                tp = doc.add_paragraph()
                tp.paragraph_format.space_after = Pt(1)
                tr = tp.add_run(job.get("title", ""))
                tr.bold = True
                tr.font.size = Pt(9.5)
                if job.get("company"):
                    tp.add_run(f"  —  {job['company']}").font.size = Pt(9.5)
                if job.get("duration"):
                    dp = doc.add_paragraph(job["duration"])
                    dp.paragraph_format.space_after = Pt(2)
                    for r in dp.runs:
                        r.italic = True
                        r.font.size = Pt(9)
                        r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
                for bullet in job.get("bullets", []):
                    bp = doc.add_paragraph(style="List Bullet")
                    bp.paragraph_format.left_indent = Inches(0.25)
                    bp.paragraph_format.space_after = Pt(1)
                    bp.add_run(bullet).font.size = Pt(9.5)
                doc.add_paragraph().paragraph_format.space_after = Pt(2)

        # Education
        education = data.get("education", [])
        if education:
            self._docx_heading(doc, "EDUCATION")
            for edu in education:
                ep = doc.add_paragraph()
                ep.paragraph_format.space_after = Pt(3)
                dr = ep.add_run(edu.get("degree", ""))
                dr.bold = True
                dr.font.size = Pt(9.5)
                suffix = f"  —  {edu.get('institution', '')}"
                if edu.get("year"):
                    suffix += f"  ({edu['year']})"
                ep.add_run(suffix).font.size = Pt(9.5)

        # Skills
        skills = data.get("skills", [])
        if skills:
            self._docx_heading(doc, "SKILLS")
            sp = doc.add_paragraph("  •  ".join(skills))
            sp.paragraph_format.space_after = Pt(4)
            for r in sp.runs:
                r.font.size = Pt(9.5)

        doc.save(buffer)
        return buffer.getvalue()

    # ── DOCX helpers ──────────────────────────────────────────────────────────

    @staticmethod
    def _docx_rule(doc: DocxDocument) -> None:
        p   = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after  = Pt(4)
        pPr = p._p.get_or_add_pPr()
        pBdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"),   "single")
        bottom.set(qn("w:sz"),    "6")
        bottom.set(qn("w:space"), "1")
        bottom.set(qn("w:color"), "3D5A80")
        pBdr.append(bottom)
        pPr.append(pBdr)

    @classmethod
    def _docx_heading(cls, doc: DocxDocument, text: str) -> None:
        p   = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after  = Pt(2)
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)
        cls._docx_rule(doc)
